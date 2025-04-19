import streamlit as st
import os
import json
from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE
import mammoth
import re
import tempfile
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
from io import BytesIO

# Function to replace text in paragraphs and runs
def replace_text_in_paragraph(paragraph, placeholder, new_text):
    # Pattern to match both [placeholder] and [Insert Type: placeholder] formats
    pattern = rf'\[(?:.*?:)?\s*{re.escape(placeholder)}\]'
    if re.search(pattern, paragraph.text):
        inline = paragraph.runs
        for i in range(len(inline)):
            if re.search(pattern, inline[i].text):
                text = re.sub(pattern, new_text, inline[i].text)
                inline[i].text = text

# Function to generate the document
def generate_document(selected_contract, form_details, local_file_path):
    try:
        doc = Document(local_file_path)
    except Exception as e:
        st.error(f"Error opening document: {e}")
        return None

    # Create styles if they don't exist
    styles = doc.styles
    if 'Heading 1' not in styles:
        styles.add_style('Heading 1', WD_STYLE_TYPE.PARAGRAPH)
        styles['Heading 1'].font.size = Pt(16)
        styles['Heading 1'].font.bold = True
    if 'Heading 2' not in styles:
        styles.add_style('Heading 2', WD_STYLE_TYPE.PARAGRAPH)
        styles['Heading 2'].font.size = Pt(14)
        styles['Heading 2'].font.bold = True

    # Apply formatting and replacements
    for paragraph in doc.paragraphs:
        # Apply heading styles
        if paragraph.text.startswith(selected_contract.upper()):
            paragraph.style = styles['Heading 1']
        elif paragraph.text.strip().startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10.", "11.", "12.")):
            paragraph.style = styles['Heading 2']
        
        # Replace placeholders
        for placeholder, value in form_details.items():
            replace_text_in_paragraph(paragraph, placeholder, value)

    return doc

# Function to convert DOCX to PDF using reportlab
def convert_docx_to_pdf(docx_content):
    # Convert DOCX to HTML
    result = mammoth.convert_to_html(docx_content)
    html_content = result.value

    # Create a PDF buffer
    buffer = BytesIO()

    # Create the PDF document
    doc = SimpleDocTemplate(buffer, pagesize=letter)

    # Create styles
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Justify', alignment=TA_JUSTIFY))
    styles.add(ParagraphStyle(name='Center', alignment=TA_CENTER))

    # Create a list to hold the flowables
    flowables = []

    # Split the HTML content into paragraphs
    paragraphs = re.split('</p>|</h1>|</h2>', html_content)

    for p in paragraphs:
        if p.strip():
            if p.startswith('<h1'):
                p = p.replace('<h1>', '').replace('</h1>', '')
                flowables.append(Paragraph(p, styles['Heading1']))
            elif p.startswith('<h2'):
                p = p.replace('<h2>', '').replace('</h2>', '')
                flowables.append(Paragraph(p, styles['Heading2']))
            else:
                p = p.replace('<p>', '').replace('</p>', '')
                flowables.append(Paragraph(p, styles['Justify']))
            flowables.append(Spacer(1, 12))

    # Build the PDF
    doc.build(flowables)

    # Get the value of the BytesIO buffer
    pdf_bytes = buffer.getvalue()
    buffer.close()

    return pdf_bytes

# CSS for styling the preview
preview_css = """
<style>
    .contract-preview {
        background-color: white;
        color: black;
        font-family: Arial, sans-serif;
        padding: 20px;
        border: 1px solid #ddd;
        border-radius: 5px;
    }
    .contract-preview h1 {
        color: #2c3e50;
        font-size: 24px;
        margin-bottom: 20px;
    }
    .contract-preview h2 {
        color: #34495e;
        font-size: 20px;
        margin-top: 15px;
        margin-bottom: 10px;
    }
    .contract-preview p {
        margin-bottom: 10px;
        line-height: 1.5;
    }
</style>
"""

def main():
    # Get the absolute path of the current script
    current_dir = os.path.dirname(os.path.abspath(__file__))

    # Set the docs directory path
    docs_dir = os.path.join(current_dir, 'docs')
    os.makedirs(docs_dir, exist_ok=True)

    # Load contract types from JSON file
    contract_types_path = os.path.join(current_dir, 'contract_types.json')
    try:
        with open(contract_types_path, 'r') as f:
            contract_types = json.load(f)
    except FileNotFoundError:
        st.error(f"Contract types file not found: {contract_types_path}")
        return
    except json.JSONDecodeError:
        st.error(f"Error parsing contract types JSON file: {contract_types_path}")
        return

    # Load placeholder questions from JSON file
    placeholder_questions_path = os.path.join(current_dir, 'placeholder_questions.json')
    try:
        with open(placeholder_questions_path, 'r') as f:
            placeholder_questions = json.load(f)
    except FileNotFoundError:
        st.error(f"Placeholder questions file not found: {placeholder_questions_path}")
        return
    except json.JSONDecodeError:
        st.error(f"Error parsing placeholder questions JSON file: {placeholder_questions_path}")
        return

    # Sidebar: Contract Types
    st.sidebar.title("Corporate & Business Contracts")
    selected_contract = st.sidebar.selectbox("Select a Contract Type", options=list(contract_types.keys()))

    # Main Content: Display Contract Form
    if selected_contract:
        st.header(f"{selected_contract} Generator")
        st.subheader("Please fill in the following details:")

        # Get the document template path
        template_filename = contract_types.get(selected_contract)
        if not template_filename:
            st.error(f"Template filename not found for {selected_contract}")
            return

        local_file_path = os.path.join(docs_dir, template_filename)

        if not os.path.exists(local_file_path):
            st.error(f"Template file not found: {local_file_path}")
            return

        # Create input fields for each placeholder
        form_details = {}
        for placeholder, question in placeholder_questions.get(selected_contract, {}).items():
            form_details[placeholder] = st.text_input(question)

        # Button to generate preview and enable downloads
        if st.button("Generate Contract"):
            st.write("Generating contract...")

            # Generate the document
            doc = generate_document(selected_contract, form_details, local_file_path)
            if doc is None:
                return

            # Save the document to a temporary file
            with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
                doc.save(tmp_file.name)
                tmp_file_path = tmp_file.name

            # Convert the document to HTML for preview
            try:
                with open(tmp_file_path, 'rb') as docx_file:
                    result = mammoth.convert_to_html(docx_file)
                    html_content = result.value
            except Exception as e:
                st.error(f"Error converting document to HTML: {e}")
                return

            # Wrap the HTML content with our custom CSS
            styled_html = f"{preview_css}<div class='contract-preview'>{html_content}</div>"

            # Display the preview
            st.subheader("Contract Preview")
            st.components.v1.html(styled_html, height=600, scrolling=True)

            # Save as DOCX
            output_docx_path = os.path.join(docs_dir, f'{selected_contract.lower().replace(" ", "_")}_output.docx')
            doc.save(output_docx_path)

            # Convert to PDF
            with open(tmp_file_path, 'rb') as docx_file:
                pdf_bytes = convert_docx_to_pdf(docx_file)

            # Download buttons
            st.subheader("Download Options")
            
            # DOCX download
            with open(output_docx_path, 'rb') as f:
                docx_bytes = f.read()
            
            st.download_button(
                label="Download Contract (DOCX)",
                data=docx_bytes,
                file_name=f'{selected_contract.lower().replace(" ", "_")}.docx',
                mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )

            # PDF download
            st.download_button(
                label="Download Contract (PDF)",
                data=pdf_bytes,
                file_name=f'{selected_contract.lower().replace(" ", "_")}.pdf',
                mime='application/pdf'
            )

            # Remove the temporary file
            os.unlink(tmp_file_path)

        # Display legal compliance notice
        st.info("This contract generator is designed to comply with Indian laws. However, it is recommended to have the final contract reviewed by a legal professional to ensure full compliance with current regulations and your specific business needs.")

if __name__ == "__main__":
    main()
